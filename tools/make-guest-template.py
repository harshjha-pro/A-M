#!/usr/bin/env python3
"""Builds web/public/templates/AM_Guest_List_Template.xlsx (FEATURES A8).

Sheets: Read me (plain steps), Guests (the list, with dropdowns), Summary (live counts).
Phone columns are text so Excel keeps every digit. Row 2 is an EXAMPLE row, skipped on import.
Run: python3 tools/make-guest-template.py   (needs openpyxl; the file is committed)
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EVENTS = ['Engagement', 'Haldi', 'Mehndi', 'Sangeet', 'Mayra', 'Wedding', 'Reception']  # the 7 seeded events
COLUMNS = ['Family name*', 'Phone', 'Other phone', 'Side*', 'Group', 'Relation', 'Area', 'City', 'Address',
           'Adults', 'Children', 'Food', 'Jain people', 'Important', *EVENTS, 'Notes']
ROWS = 2000
OUT = Path(__file__).resolve().parent.parent / 'web/public/templates/AM_Guest_List_Template.xlsx'

wb = Workbook()
readme = wb.active
readme.title = 'Read me'
lines = [
    ('A&M Wedding — guest list template', True),
    ('', False),
    ('1. Fill the "Guests" sheet: one row per family (not per person).', False),
    ('2. Family name and Side are needed. Everything else is optional.', False),
    ('3. Phone: type the 10-digit mobile number, like 98290 12345. The column keeps every digit.', False),
    ('4. Adults and Children: how many people will come from this family. Empty = 2 adults, 0 children.', False),
    ('5. Food: Veg, Jain or Mixed (a veg family with some Jain members — then fill "Jain people").', False),
    ('6. For each function, choose Yes if this family is invited to it.', False),
    ('7. Row 2 is an EXAMPLE. Leave it or delete it: rows whose name starts with EXAMPLE are skipped.', False),
    ('8. In the app: Guests → Import a list → choose this file. You will see a preview before anything is saved.', False),
    ('', False),
    ('The Summary sheet counts families and people as you type.', False),
]
for i, (text, bold) in enumerate(lines, start=1):
    readme.cell(row=i, column=1, value=text).font = Font(bold=bold, size=14 if bold else 12)
readme.column_dimensions['A'].width = 110

g = wb.create_sheet('Guests')
head_fill = PatternFill('solid', fgColor='F3E3E6')
for c, name in enumerate(COLUMNS, start=1):
    cell = g.cell(row=1, column=c, value=name)
    cell.font = Font(bold=True)
    cell.fill = head_fill
    cell.alignment = Alignment(wrap_text=True, vertical='center')
    g.column_dimensions[get_column_letter(c)].width = 26 if c in (1, 9) else (16 if c in (2, 3, 5, 6, 7, 8) else 12)
g.freeze_panes = 'B2'
example = ['EXAMPLE Ramesh Sharma & family', '98290 12345', '', "Groom's side", 'Papa office', 'Friend', 'Shastri Nagar', 'Bhilwara', '',
           3, 1, 'Mixed', 2, 'No', 'No', 'No', 'Yes', 'No', 'No', 'Yes', 'Yes', 'Arriving late']
for c, v in enumerate(example, start=1):
    g.cell(row=2, column=c, value=v if v != '' else None).font = Font(italic=True, color='7A6A60')

col = {name: get_column_letter(i) for i, name in enumerate(COLUMNS, start=1)}
for name in ('Phone', 'Other phone'):  # text, so Excel keeps leading digits and never shows 9.83E+09
    for r in range(2, ROWS + 2):
        g[f'{col[name]}{r}'].number_format = '@'

def dropdown(values, column, prompt):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True, showDropDown=False)
    dv.promptTitle, dv.prompt = column.rstrip('*'), prompt
    dv.showInputMessage = True
    g.add_data_validation(dv)
    dv.add(f'{col[column]}2:{col[column]}{ROWS + 1}')

dropdown(["Bride's side", "Groom's side", 'Both sides'], 'Side*', 'Whose side invited this family')
dropdown(['Veg', 'Jain', 'Mixed'], 'Food', 'Mixed = veg family with some Jain members')
dropdown(['Yes', 'No'], 'Important', 'Shown with a star in the app')
for e in EVENTS:
    dropdown(['Yes', 'No'], e, f'Invited to {e}?')
nums = DataValidation(type='whole', operator='between', formula1='0', formula2='50', allow_blank=True)
nums.error, nums.showErrorMessage = 'Use a number from 0 to 50.', True
g.add_data_validation(nums)
for name in ('Adults', 'Children', 'Jain people'):
    nums.add(f'{col[name]}2:{col[name]}{ROWS + 1}')

s = wb.create_sheet('Summary')
rng = lambda name: f"Guests!${col[name]}$3:${col[name]}${ROWS + 1}"  # row 2 is the example
s['A1'] = 'Summary (counts update as you type; the EXAMPLE row is not counted)'
s['A1'].font = Font(bold=True, size=14)
rows = [('Families', f'=COUNTA({rng("Family name*")})'),
        ('People (adults + children; empty adults count as 2)',
         f'=SUM({rng("Adults")})+SUM({rng("Children")})+2*COUNTIFS({rng("Family name*")},"<>",{rng("Adults")},"")'),
        ("Bride's side families", f'=COUNTIF({rng("Side*")},"Bride\'s side")'),
        ("Groom's side families", f'=COUNTIF({rng("Side*")},"Groom\'s side")'),
        ('Both sides families', f'=COUNTIF({rng("Side*")},"Both sides")')]
rows += [(f'{e}: families invited', f'=COUNTIF({rng(e)},"Yes")') for e in EVENTS]
for i, (label, formula) in enumerate(rows, start=3):
    s.cell(row=i, column=1, value=label)
    s.cell(row=i, column=2, value=formula)
s.column_dimensions['A'].width = 52
s.column_dimensions['B'].width = 14
wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'Wrote {OUT}')
